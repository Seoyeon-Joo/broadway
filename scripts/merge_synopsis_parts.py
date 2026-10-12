#!/usr/bin/env python3
"""샤드별 part_*.xlsx 를 데이터셋별(broadway/hollywood)로 합쳐 data/synopses_*.xlsx 로 저장."""
import glob
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def merge(name):
    files = sorted(glob.glob(f"parts/part_{name}_*.xlsx"))
    syn = pd.concat([pd.read_excel(f, "synopses") for f in files], ignore_index=True)
    full = pd.concat([pd.read_excel(f, "full_plot") for f in files], ignore_index=True)
    syn = syn.sort_values("work_id").reset_index(drop=True)
    full = full.sort_values("work_id").reset_index(drop=True)
    out = f"data/synopses_{name}.xlsx"
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        syn.to_excel(w, sheet_name="synopses", index=False)
        full.to_excel(w, sheet_name="full_plot", index=False)
        ws = w.sheets["synopses"]
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="305496")
        for col, wd in zip("ABCDEFGHIJKL", [8, 34, 9, 7, 90, 50, 18, 12, 12, 11, 36, 30]):
            ws.column_dimensions[col].width = wd
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"
        w.sheets["full_plot"].column_dimensions["C"].width = 120
    print(out, len(syn), "works;", int((syn["needs_review"] == "Y").sum()), "need review")


for n in ("broadway_panel", "hollywood"):
    merge(n)
