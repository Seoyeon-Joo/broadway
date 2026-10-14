#!/usr/bin/env python3
"""
Wikipedia 줄거리를 못 찾은 Broadway 작품을 더 넓은 규칙으로 다시 찾는다.

collect_synopses.py 와 같은 MediaWiki API(영어 Wikipedia)를 쓰되 다음을 추가한다.
 1) 제목이 Plot/Synopsis 가 아니어도 인식: "Act I/Act One" 소제목, Overview/Scenes 등 (하위 소제목 포함)
 2) 소제목이 아닌 본문 줄 "ACT ONE", "ACT TWO" 로만 나뉜 줄거리 (예: New York, New York 페이지)
 3) 후보 페이지를 직접 지정(공연 페이지 -> 원작 페이지 순). 원작(소설/영화) 줄거리는 kind='source material' 로 구분
 4) 못 찾으면 그 페이지의 소제목 목록을 headings_seen 열에 남김 (규칙 보완용)
출력: data/synopses_retry_broadway.xlsx  (본 파일에는 자동으로 합치지 않는다)
본문은 CC BY-SA 4.0 이므로 출처 URL 표기가 필요하다.
"""
import argparse
import os
import re
import sys
import time

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from collect_synopses import api, get_text, split_sections  # noqa: E402

OUT = "data/synopses_retry_broadway.xlsx"
MAX_CHARS = 32000

# work_id: (title, year, [공연 페이지 후보], [원작 페이지 후보])
TARGETS = {
    7: ("Lackawanna Blues", 2021, ["Lackawanna Blues", "Lackawanna Blues (play)"], ["Lackawanna Blues (film)"]),
    8: ("SIX: The Musical", 2021, ["Six (musical)"], []),
    11: ("Chicken & Biscuits", 2021, ["Chicken & Biscuits (play)", "Chicken & Biscuits"], []),
    13: ("Is This A Room", 2021, ["Is This a Room", "Is This a Room (play)"], []),
    23: ("Ain't Too Proud", 2021, ["Ain't Too Proud"], []),
    28: ("Trouble in Mind", 2021, ["Trouble in Mind (play)"], []),
    48: ("Mr. Saturday Night", 2022, ["Mr. Saturday Night (musical)"], ["Mr. Saturday Night"]),
    49: ("The Little Prince", 2022, ["The Little Prince (musical)"], ["The Little Prince"]),
    58: ("The Kite Runner", 2022, ["The Kite Runner (play)"], ["The Kite Runner"]),
    65: ("Almost Famous", 2022, ["Almost Famous (musical)"], ["Almost Famous"]),
    68: ("Walking with Ghosts", 2022, ["Walking with Ghosts"], []),
    72: ("A Beautiful Noise", 2022, ["A Beautiful Noise (musical)"], []),
    86: ("Life of Pi", 2023, ["Life of Pi (play)"], ["Life of Pi"]),
    91: ("New York, New York", 2023, ["New York, New York (musical)"], ["New York, New York (1977 film)"]),
    127: ("The Heart of Rock and Roll", 2024, ["The Heart of Rock and Roll (musical)"], []),
    132: ("Patriots", 2024, ["Patriots (play)"], []),
    136: ("Home", 2024, ["Home (Williams play)", "Home (play)"], []),
    139: ("JOB", 2024, ["Job (play)", "JOB (play)"], []),
    149: ("A Wonderful World", 2024, ["A Wonderful World (musical)"], []),
    151: ("Tammy Faye", 2024, ["Tammy Faye (musical)"], []),
    155: ("Cult of Love", 2024, ["Cult of Love"], []),
    169: ("The Picture of Dorian Gray", 2025, ["The Picture of Dorian Gray (play)"], ["The Picture of Dorian Gray"]),
    188: ("Little Bear Ridge Road", 2025, ["Little Bear Ridge Road"], []),
    199: ("Dog Day Afternoon", 2026, ["Dog Day Afternoon (play)"], ["Dog Day Afternoon"]),
    204: ("Beaches", 2026, ["Beaches (musical)"], ["Beaches (1988 film)"]),
    208: ("Titaníque", 2026, ["Titaníque", "Titanique"], []),
}

ACT_HEAD = re.compile(r"^(act|part)\s+([ivx]+|one|two|three|\d+)\b", re.I)
PLOT_ANY = re.compile(r"\b(plot|synopsis|summary|story|storyline|premise|overview|scenes?)\b", re.I)
BAD_HEAD = re.compile(r"critic|reception|adaptation|production|soundtrack|musical numbers|cast|characters|roles|recording|award|history|"
                      r"development|background|references|notes|external|legacy|revival|casting|awards", re.I)
ACT_LINE = re.compile(r"^\s*(?:ACT|Act)\s+(?:ONE|TWO|THREE|I{1,3}|1|2|3)\s*[.:]?\s*$")
STOP_HEAD = re.compile(r"^(musical numbers|cast|characters|roles|productions?|recordings?|awards?|reception|critical|references|notes|"
                       r"external links|see also|original cast|principal roles|casting)", re.I)


def words(s):
    return len(s.split())


def prose_words(block):
    return sum(words(ln) for ln in block.split("\n") if len(ln) > 100)


def plot_from_headings(text):
    secs = split_sections(text)
    for i, (lvl, head, body) in enumerate(secs):
        if lvl < 2 or not head:
            continue
        direct = lvl == 2 and PLOT_ANY.fullmatch(head.strip().lower().split(" of the")[0] or "")
        if not (ACT_HEAD.match(head) or PLOT_ANY.search(head)) or BAD_HEAD.search(head):
            continue
        chunks, labels = [body], [head]
        for lvl2, head2, body2 in secs[i + 1:]:
            deeper = lvl2 > lvl
            sibling_act = lvl2 == lvl and ACT_HEAD.match(head) and ACT_HEAD.match(head2)
            if not (deeper or sibling_act):
                break
            if head2:
                chunks.append(head2 + ".")
            chunks.append(body2)
            labels.append(head2)
        txt = "\n\n".join(c for c in chunks if c)
        if prose_words(txt) >= 120 or words(txt) >= 200:
            return txt, "heading: " + head + (f" (+{len(labels) - 1} sub)" if len(labels) > 1 else "")
    return None, None


def plot_from_act_lines(text):
    """소제목 없이 'ACT ONE' 같은 줄로만 나뉜 본문. 가장 줄글이 많은 블록을 고른다."""
    lines = text.split("\n")
    idx = [i for i, ln in enumerate(lines) if ACT_LINE.match(ln)]
    best, best_pw = None, 0
    for k in idx:
        end = len(lines)
        for j in range(k + 1, len(lines)):
            m = re.match(r"^(={2,6})\s*(.+?)\s*\1\s*$", lines[j])
            if m and STOP_HEAD.match(m.group(2)):
                end = j
                break
        block = "\n".join(lines[k:end]).strip()
        pw = prose_words(block)
        if pw > best_pw:
            best, best_pw = block, pw
    if best and best_pw >= 150:
        return best, "inline ACT markers"
    return None, None


def extract(text):
    from collect_synopses import plot_section
    p, h = plot_section(text)
    if p and words(p) >= 120:
        return p, "heading: " + h
    p, lab = plot_from_headings(text)
    if p:
        return p, lab
    return plot_from_act_lines(text)


def headings(text):
    return "; ".join(f"{'#' * l}{h}" for l, h, _ in split_sections(text) if h)[:600]


def try_pages(titles, kind, dbg, seen):
    for t in titles:
        real, text = get_text(t)
        time.sleep(0.4)
        if not text:
            dbg.append(f"{t}: missing")
            continue
        plot, lab = extract(text)
        dbg.append(f"{t}: {'plot via ' + lab if plot else 'no plot'} ({words(text)}w)")
        if plot:
            return plot[:MAX_CHARS], lab, real, kind
        seen.append(f"[{real}] {headings(text)}")
    return None, None, None, None


def process(wid, spec):
    title, year, stage, source = spec
    dbg, seen = [], []
    plot, lab, page, kind = try_pages(stage, "stage page", dbg, seen)
    if not plot and source:
        plot, lab, page, kind = try_pages(source, "source material (not the stage adaptation)", dbg, seen)
    status = "ok" if plot else "not found"
    return {
        "work_id": wid, "title": title, "year": year, "kind": kind or "",
        "page_title": page or "", "url": ("https://en.wikipedia.org/wiki/" + page.replace(" ", "_")) if page else "",
        "section_used": lab or "", "plot_words": words(plot) if plot else 0, "synopsis": plot or "",
        "status": status, "headings_seen": " | ".join(seen)[:1500], "debug": "; ".join(dbg)[:800],
    }


def save(rows, path):
    cols = ["work_id", "title", "year", "kind", "page_title", "url", "section_used", "plot_words", "synopsis", "status",
            "headings_seen", "debug"]
    wb = Workbook()
    ws = wb.active
    ws.title = "retry"
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for col, w in zip("ABCDEFGHIJKL", [8, 30, 7, 22, 30, 45, 26, 10, 90, 11, 60, 60]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=OUT)
    a = ap.parse_args()
    rows = []
    for n, (wid, spec) in enumerate(TARGETS.items(), 1):
        try:
            rec = process(wid, spec)
        except Exception as e:
            rec = {"work_id": wid, "title": spec[0], "year": spec[1], "kind": "", "page_title": "", "url": "", "section_used": "",
                   "plot_words": 0, "synopsis": "", "status": "error", "headings_seen": "", "debug": str(e)[:300]}
        rows.append(rec)
        print(f"[{n}/{len(TARGETS)}] {rec['title']} -> {rec['status']} | {rec['kind']} | {rec['section_used']} | {rec['debug'][:140]}", flush=True)
    save(rows, a.out)
    print("saved", a.out, "ok:", sum(1 for r in rows if r["status"] == "ok"), "of", len(rows))


if __name__ == "__main__":
    sys.exit(main())
